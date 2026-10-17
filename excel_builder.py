import io
import os
import re
import difflib
from typing import List, Dict, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def sanitize_sheet_name(name: str, existing_names: List[str] = None) -> str:
    """Sanitize sheet name for Excel (max 31 chars, no invalid chars)."""
    if not name or not name.strip():
        name = "Unassigned"
    # Remove invalid characters: \ / ? * [ ] :
    clean = re.sub(r'[\\/\?\*\[\]:]', '_', name.strip())
    # Excel sheet title limit is 31 characters
    clean = clean[:31].strip()
    if not clean:
        clean = "Sheet"
    
    if existing_names is None:
        return clean
        
    final_name = clean
    counter = 1
    while final_name.lower() in [n.lower() for n in existing_names]:
        suffix = f"_{counter}"
        max_base_len = 31 - len(suffix)
        final_name = clean[:max_base_len] + suffix
        counter += 1
    return final_name


def format_excel_sheet(ws, title: str, rows: List[Dict[str, Any]]):
    """Format an openpyxl worksheet with professional styling, formulas, and auto column widths."""
    
    headers = [
        ("File Name", 24, Alignment(horizontal="left", vertical="center")),
        ("Customer Name", 26, Alignment(horizontal="left", vertical="center")),
        ("Phone Number", 15, Alignment(horizontal="center", vertical="center")),
        ("Digital Address", 18, Alignment(horizontal="left", vertical="center")),
        ("Account Number", 16, Alignment(horizontal="center", vertical="center")),
        ("Meter Number", 16, Alignment(horizontal="center", vertical="center")),
        ("Bill Period", 30, Alignment(horizontal="left", vertical="center")),
        ("Bill Number", 20, Alignment(horizontal="center", vertical="center")),
        ("Issuing Center", 24, Alignment(horizontal="left", vertical="center")),
        ("Issue Date", 14, Alignment(horizontal="center", vertical="center")),
        ("Due Date", 14, Alignment(horizontal="center", vertical="center")),
        ("Previous Reading", 17, Alignment(horizontal="right", vertical="center")),
        ("Current Reading", 17, Alignment(horizontal="right", vertical="center")),
        ("Consumption (kWh)", 18, Alignment(horizontal="right", vertical="center")),
        ("Current Bill (GHS)", 18, Alignment(horizontal="right", vertical="center")),
        ("Previous Balance (GHS)", 22, Alignment(horizontal="right", vertical="center")),
        ("Amount Payable (GHS)", 22, Alignment(horizontal="right", vertical="center")),
    ]

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    
    data_font = Font(name="Segoe UI", size=10)
    total_font = Font(name="Segoe UI", size=11, bold=True, color="000000")
    total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    data_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    total_top = Side(border_style="thin", color="1F4E79")
    total_bottom = Side(border_style="double", color="1F4E79")
    total_border = Border(top=total_top, bottom=total_bottom, left=thin_border_side, right=thin_border_side)

    # Write Header
    ws.row_dimensions[1].height = 28
    for col_idx, (header_label, width, align) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Write Data Rows
    current_row = 2
    for item in rows:
        ws.row_dimensions[current_row].height = 20
        # Alternating background
        row_fill = PatternFill(start_color="F9FAFB" if current_row % 2 == 0 else "FFFFFF",
                               end_color="F9FAFB" if current_row % 2 == 0 else "FFFFFF",
                               fill_type="solid")

        row_data = [
            (str(item.get("File Name") or ""), "@", headers[0][2]),
            (str(item.get("Customer Name") or ""), "@", headers[1][2]),
            (str(item.get("Phone Number") or ""), "@", headers[2][2]),
            (str(item.get("Digital Address") or ""), "@", headers[3][2]),
            (str(item.get("Account Number") or ""), "@", headers[4][2]),
            (str(item.get("Meter Number") or ""), "@", headers[5][2]),
            (str(item.get("Bill Period") or ""), "@", headers[6][2]),
            (str(item.get("Bill Number") or ""), "@", headers[7][2]),
            (str(item.get("Issuing Center") or ""), "@", headers[8][2]),
            (str(item.get("Issue Date") or ""), "@", headers[9][2]),
            (str(item.get("Due Date") or ""), "@", headers[10][2]),
            (item.get("Previous Reading"), "#,##0", headers[11][2]),
            (item.get("Current Reading"), "#,##0", headers[12][2]),
            (item.get("Consumption"), "#,##0", headers[13][2]),
            (item.get("Current Bill"), "#,##0.00;[Red]-#,##0.00;0.00", headers[14][2]),
            (item.get("Previous Balance"), "#,##0.00;[Red]-#,##0.00;0.00", headers[15][2]),
            (item.get("Amount Payable"), "#,##0.00;[Red]-#,##0.00;0.00", headers[16][2]),
        ]

        for col_idx, (val, num_format, align) in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = data_font
            cell.fill = row_fill
            cell.alignment = align
            cell.border = data_border
            if num_format:
                cell.number_format = num_format

        current_row += 1

    # Write Total / Summary Row if rows exist
    if rows:
        ws.row_dimensions[current_row].height = 24
        # Label in Due Date column (col 11) or Bill Number
        for col_idx in range(1, 14):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = total_border
            cell.alignment = Alignment(horizontal="right", vertical="center")
            if col_idx == 11:
                cell.value = "TOTAL"
                cell.alignment = Alignment(horizontal="center", vertical="center")

        # Sum formulas for Consumption (14), Current Bill (15), Prev Balance (16), Amount Payable (17)
        numeric_sum_cols = [
            (14, "#,##0"),
            (15, "#,##0.00;[Red]-#,##0.00;0.00"),
            (16, "#,##0.00;[Red]-#,##0.00;0.00"),
            (17, "#,##0.00;[Red]-#,##0.00;0.00")
        ]
        
        start_row = 2
        end_row = current_row - 1
        for col_idx, num_fmt in numeric_sum_cols:
            col_letter = get_column_letter(col_idx)
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = f"=SUM({col_letter}{start_row}:{col_letter}{end_row})"
            cell.font = total_font
            cell.fill = total_fill
            cell.border = total_border
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = num_fmt

    # Adjust column widths based on content + header with minimums
    for col_idx, (header_label, min_width, align) in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(header_label)
        for r in range(2, current_row):
            val = ws.cell(row=r, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 4, 45))

    # Freeze top row
    ws.freeze_panes = "A2"


def find_similar_cluster(key: str, existing_keys: List[str], threshold: float = 0.82) -> Optional[str]:
    """Find an existing cluster key that is similar or closely matching."""
    key_clean = key.strip().lower()
    for cand in existing_keys:
        cand_clean = cand.strip().lower()
        if key_clean == cand_clean:
            return cand
        # Substring / abbreviation match
        if (len(key_clean) >= 4 and key_clean in cand_clean) or (len(cand_clean) >= 4 and cand_clean in key_clean):
            return cand
        # Fuzzy ratio match
        ratio = difflib.SequenceMatcher(None, key_clean, cand_clean).ratio()
        if ratio >= threshold:
            return cand
    return None


def generate_excel_workbook(invoices: List[Dict[str, Any]], group_by: str = "center") -> Workbook:
    """Generate an Excel workbook with sheets grouped by Issuing Center (or Customer/Both) with similarity clustering."""
    wb = Workbook()
    default_sheet = wb.active

    # Group invoices with fuzzy similarity clustering
    grouped: Dict[str, List[Dict[str, Any]]] = {}
    cluster_keys: List[str] = []

    for inv in invoices:
        raw_center = (inv.get("Issuing Center") or "Unknown District").strip().title()
        raw_customer = (inv.get("Customer Name") or "Unknown Customer").strip()

        if group_by == "customer":
            raw_key = raw_customer
        elif group_by == "both":
            raw_key = f"{raw_customer} - {raw_center}"
        else: # default "center"
            raw_key = raw_center

        if not raw_key:
            raw_key = "General"

        # Check if there is an existing similar cluster
        matched_cluster = find_similar_cluster(raw_key, cluster_keys)
        if matched_cluster:
            cluster_name = matched_cluster
        else:
            cluster_name = raw_key
            cluster_keys.append(cluster_name)
            grouped[cluster_name] = []

        grouped[cluster_name].append(inv)

    sheet_names_created = []

    # If multiple clusters exist or multiple bills, add an "All Bills Summary" sheet first
    if len(grouped) > 1 or len(invoices) > 1:
        ws_all = default_sheet
        ws_all.title = "All Bills Summary"
        format_excel_sheet(ws_all, "All Bills Summary", invoices)
        sheet_names_created.append(ws_all.title)
        first_created = True
    else:
        first_created = False

    # Create sheets for each cluster
    for cluster_name, cluster_invoices in grouped.items():
        sheet_title = sanitize_sheet_name(cluster_name, sheet_names_created)
        if not first_created:
            ws = default_sheet
            ws.title = sheet_title
            first_created = True
        else:
            ws = wb.create_sheet(title=sheet_title)
        
        sheet_names_created.append(sheet_title)
        format_excel_sheet(ws, sheet_title, cluster_invoices)

    return wb


def generate_excel_bytes(invoices: List[Dict[str, Any]], group_by: str = "center") -> bytes:
    """Generate Excel file in-memory as bytes with specified grouping."""
    wb = generate_excel_workbook(invoices, group_by=group_by)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


if __name__ == "__main__":
    from extractor import extract_from_pdf_path
    sample_file = r"c:\Users\wwukpo\Desktop\ECG BILLS\ABORE HOUSE ELECTRICITY BILL.pdf"
    if os.path.exists(sample_file):
        sample_data = extract_from_pdf_path(sample_file)
        # Create 2 mock entries with same center and 1 with different center to test grouping
        mock_list = [
            sample_data,
            {**sample_data, "Customer Name": "KWAME MENSAH", "Account Number": "746624899", "Amount Payable": -120.50},
            {**sample_data, "Issuing Center": "Accra West District", "Customer Name": "AMA SERWAA", "Amount Payable": 450.00}
        ]
        wb = generate_excel_workbook(mock_list)
        test_out = r"C:\Users\wwukpo\.gemini\antigravity\scratch\ecg_invoice_extractor\test_output.xlsx"
        wb.save(test_out)
        print("Excel workbook created successfully at:", test_out)
        print("Sheets in workbook:", wb.sheetnames)
